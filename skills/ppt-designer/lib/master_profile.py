"""Lê qualquer slide mestre (.pptx/.potx) e grava um profile.json que o slide_style.py usa no lugar de constantes fixas.

O profile traz:
- tamanho do slide;
- fontes e cores do tema;
- paleta semântica sugerida;
- para cada layout: os placeholders com caixa em polegadas e as formas decorativas (zonas proibidas);
- a grade inferida: margens, colunas, linhas de cabeçalho, título, subtítulo e rodapé.

Uso:
    python3 lib/master_profile.py <mestre.pptx> <saida_dir>
Gera <saida_dir>/profile.json e <saida_dir>/layouts_preview.pptx (um slide por layout, com cada placeholder
rotulado). Exporte esse preview para PDF/PNG e OLHE antes de desenhar: o profile é inferido e pode errar.
"""
import colorsys
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

EMU = 914400
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def inch(v):
    return round((v or 0) / EMU, 3)


def box(sh):
    return [inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)]


# ---------------------------------------------------------------- tema
def theme(prs):
    master = prs.slide_masters[0]
    part = next(r.target_part for r in master.part.rels.values() if r.reltype.endswith("/theme"))
    from lxml import etree
    root = etree.fromstring(part.blob)
    colors = {}
    scheme = root.find(f".//{{{A}}}clrScheme")
    for el in scheme:
        name = el.tag.split("}")[1]
        c = el[0]
        colors[name] = (c.get("val") if c.tag.endswith("srgbClr") else c.get("lastClr")) or "000000"
    fonts = {}
    for kind in ("majorFont", "minorFont"):
        lat = root.find(f".//{{{A}}}{kind}/{{{A}}}latin")
        fonts[kind] = lat.get("typeface") if lat is not None else "Calibri"
    return colors, fonts


def _hls(h):
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return colorsys.rgb_to_hls(r, g, b)


def _hex(r, g, b):
    return "".join(f"{round(v * 255):02X}" for v in (r, g, b))


def tint(h, k):
    """Mistura com branco: k=0 mantém, k=1 vira branco."""
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return _hex(r + (1 - r) * k, g + (1 - g) * k, b + (1 - b) * k)


def shade(h, k):
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return _hex(r * (1 - k), g * (1 - k), b * (1 - k))


def palette(colors):
    """Paleta semântica sugerida. Os papéis são fixos; os hex vêm do tema e podem ser sobrescritos em deck_prefs.json.
    FOCUS = a cor da empresa-foco: o acento mais escuro e saturado do tema (dk2 se for colorido).
    ACCENT = o único destaque: o acento mais saturado e claro que se distingue do FOCUS."""
    cands = [colors.get(k) for k in ("dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if colors.get(k)]
    sat = lambda h: _hls(h)[2]
    lum = lambda h: _hls(h)[1]
    colored = [c for c in cands if sat(c) > 0.25] or cands
    focus = min(colored, key=lambda c: lum(c) - 0.2 * sat(c))
    others = [c for c in colored if c != focus and abs(lum(c) - lum(focus)) > 0.12] or [tint(focus, 0.35)]
    accent = max(others, key=lambda c: sat(c) - abs(lum(c) - 0.45))
    return {
        "FOCUS": focus,                      # empresa-foco, série principal
        "FOCUS_TINT": tint(focus, 0.68),     # empresa-foco fora de destaque, "restante"
        "FOCUS_DARK": shade(focus, 0.3),     # lateral de sólidos 3D
        "FOCUS_TOP": tint(focus, 0.85),      # topo de sólidos 3D
        "ACCENT": accent,                    # ÚNICO destaque
        "PEER": "AEB8C6",                    # terceiros / concorrentes
        "PEER_DARK": "6B7A90",               # linha/pill de terceiros
        "NEG": "D62F2F",                     # só negativo / alerta
        "FG": colors.get("dk1", "0A1424") if colors.get("dk1", "000000") != "000000" else "0A1424",
        "FG2": "2A3A55", "MUTED": "5B6B85", "AXIS": "8A96A6", "WHITE": "FFFFFF",
    }


# ---------------------------------------------------------------- layouts
PH_KIND = {
    PP_PLACEHOLDER.TITLE: "title", PP_PLACEHOLDER.CENTER_TITLE: "title", PP_PLACEHOLDER.SUBTITLE: "subtitle",
    PP_PLACEHOLDER.BODY: "body", PP_PLACEHOLDER.OBJECT: "object", PP_PLACEHOLDER.CHART: "chart",
    PP_PLACEHOLDER.TABLE: "table", PP_PLACEHOLDER.PICTURE: "picture", PP_PLACEHOLDER.FOOTER: "footer",
    PP_PLACEHOLDER.SLIDE_NUMBER: "slide_number", PP_PLACEHOLDER.DATE: "date",
}


def _font_size(sh):
    try:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    return r.font.size.pt
        lvl = sh._element.find(".//" + qn("a:lvl1pPr") + "/" + qn("a:defRPr"))
        if lvl is not None and lvl.get("sz"):
            return int(lvl.get("sz")) / 100
    except Exception:
        pass
    return None


def decor(shapes, W, H):
    """Formas que não são placeholder (logos, curvas, faixas): viram zonas proibidas para conteúdo."""
    out = []
    for sh in shapes:
        if sh.is_placeholder:
            continue
        b = box(sh)
        if b[2] >= W * 0.95 and b[3] >= H * 0.95:      # fundo de página inteira não é zona proibida
            continue
        kind = {MSO_SHAPE_TYPE.PICTURE: "picture", MSO_SHAPE_TYPE.FREEFORM: "freeform",
                MSO_SHAPE_TYPE.GROUP: "group"}.get(sh.shape_type, "shape")
        txt = sh.text_frame.text.strip()[:40] if sh.has_text_frame else ""
        out.append({"name": sh.name, "kind": kind, "box": b, "text": txt})
    return out


def layouts(prs, W, H):
    res = []
    for i, lay in enumerate(prs.slide_layouts):
        phs = []
        for sh in lay.placeholders:
            pf = sh.placeholder_format
            phs.append({"idx": pf.idx, "kind": PH_KIND.get(pf.type, str(pf.type).split(".")[-1].split(" ")[0].lower()),
                        "name": sh.name, "box": box(sh), "font_pt": _font_size(sh)})
        res.append({"index": i, "name": lay.name, "placeholders": sorted(phs, key=lambda p: (p["box"][1], p["box"][0])),
                    "decor": decor(lay.shapes, W, H)})
    return res


def _cluster(vals, tol=0.15):
    out = []
    for v in sorted(vals):
        if out and v - out[-1][-1] <= tol:
            out[-1].append(v)
        else:
            out.append([v])
    return [round(sum(c) / len(c), 2) for c in out if len(c) >= 2] or [round(c[0], 2) for c in out]


def _on_slide(b, W, H):
    x, y, w, h = b
    return x + w > 0.05 and y + h > 0.05 and x < W - 0.05 and y < H - 0.05


def _layout_title(l, W, H):
    """Título de conteúdo do layout: placeholder title (ou, em mestres que usam 'body', a caixa larga mais alta),
    no topo da página. Títulos centralizados de capa/divisor (no meio da página) não contam."""
    cands = [p for p in l["placeholders"] if _on_slide(p["box"], W, H) and p["box"][1] < H * 0.2]
    t = [p for p in cands if p["kind"] == "title"] or \
        [p for p in cands if p["kind"] == "body" and p["box"][2] > W * 0.5]
    return min(t, key=lambda p: p["box"][1]) if t else None


def grid(lays, master_decor, W, H):
    """Grade inferida dos layouts de CONTEÚDO (título no topo + caixa de conteúdo abaixo); capa e divisores são
    ignorados. Título = a caixa de título mais comum entre eles; subtítulo = caixa larga logo abaixo; rodapé =
    placeholder no último 10% da página."""
    wide = lambda p: p["box"][2] > W * 0.5
    content_kinds = ("body", "object", "chart", "table", "picture")
    picked, titles = [], []
    for l in lays:
        t = _layout_title(l, W, H)
        if not t:
            continue
        tb = t["box"][1] + t["box"][3]
        body = [p for p in l["placeholders"] if p is not t and _on_slide(p["box"], W, H)
                and p["kind"] in content_kinds and p["box"][1] > tb and p["box"][1] < H * 0.9 and p["box"][2] > 0.8]
        if body:
            picked.append((l, t, body)); titles.append(t)
    if not picked:                                    # mestre sem layout de conteúdo: área útil genérica
        return {"slide_w": W, "slide_h": H, "content_left": 0.6, "content_right": round(W - 0.6, 3),
                "content_top": 1.4, "content_bottom": round(H - 0.7, 3), "gutter": 0.33, "column_x": [], "row_y": [],
                "title": None, "subtitle": None, "footer": None, "forbidden": [], "warning": "nenhum layout de conteúdo"}
    key = lambda p: tuple(round(v, 1) for v in p["box"])
    common = max(set(key(t) for t in titles), key=lambda k: sum(key(t) == k for t in titles))
    title = next(t for t in titles if key(t) == common)
    tb = title["box"][1] + title["box"][3]
    subs = [p for _, t, body in picked for p in body if p["kind"] in ("subtitle", "body") and wide(p)
            and -0.05 <= p["box"][1] - tb < 0.45 and p["box"][3] < 0.6]
    sub = min(subs, key=lambda p: p["box"][1]) if subs else None
    allph = [p for l in lays for p in l["placeholders"] if _on_slide(p["box"], W, H)]
    foot = sorted([p for p in allph if p["box"][1] > H * 0.9 and p["kind"] in ("footer", "body")],
                  key=lambda p: -p["box"][2])
    content = [p for _, _, body in picked for p in body if not (sub and key(p) == key(sub)) and p["box"][1] < H * 0.9]
    left = min(p["box"][0] for p in content)
    right = max(p["box"][0] + p["box"][2] for p in content)
    top = min(p["box"][1] for p in content)
    bottom = max(p["box"][1] + p["box"][3] for p in content)
    forb = []
    for d in master_decor:
        x, y, w, h = d["box"]
        if not _on_slide(d["box"], W, H) or min(w, h) < 0.02:      # fora do slide ou fio de espessura zero
            continue
        x2, y2 = min(x + w, W), min(y + h, H)
        x, y = max(x, 0), max(y, 0)
        forb.append([round(x, 3), round(y, 3), round(x2 - x, 3), round(y2 - y, 3)])
    inside = lambda a, b: a is not b and a[0] >= b[0] and a[1] >= b[1] and \
        a[0] + a[2] <= b[0] + b[2] and a[1] + a[3] <= b[1] + b[3]
    forb = [a for a in forb if not any(inside(a, b) for b in forb)]   # caixa dentro de outra: redundante
    pick = lambda p: p and {"idx": p["idx"], "box": p["box"], "font_pt": p["font_pt"]}
    return {
        "slide_w": W, "slide_h": H,
        "content_left": round(left, 3), "content_right": round(right, 3),
        "content_top": round(top, 3), "content_bottom": round(bottom, 3),
        "gutter": 0.33,
        "column_x": _cluster([p["box"][0] for p in content]),
        "row_y": _cluster([p["box"][1] for p in content]),
        "content_layouts": [l["index"] for l, _, _ in picked],
        "title": pick(title), "subtitle": pick(sub),
        "footer": foot and {"idx": foot[0]["idx"], "box": foot[0]["box"]},
        "forbidden": forb,
    }


def profile(master_path):
    prs = Presentation(master_path)
    W, H = inch(prs.slide_width), inch(prs.slide_height)
    colors, fonts = theme(prs)
    m_decor = decor(prs.slide_masters[0].shapes, W, H)
    lays = layouts(prs, W, H)
    return {
        "master": str(Path(master_path).resolve()),
        "slide": [W, H],
        "theme_colors": colors,
        "fonts": {"heading": fonts["majorFont"], "body": fonts["minorFont"]},
        "palette": palette(colors),
        "master_decor": m_decor,
        "layouts": lays,
        "grid": grid(lays, m_decor, W, H),
        "notes": "Inferido automaticamente. Confira no layouts_preview e corrija aqui o que estiver errado "
                 "(ex.: grid.forbidden, palette.FOCUS).",
    }


def preview(master_path, prof, out):
    """Um slide por layout, cada placeholder preenchido com 'idx · kind' e as zonas proibidas em vermelho."""
    prs = Presentation(master_path)
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId); ids.remove(sid)
    for lay in prof["layouts"]:
        s = prs.slides.add_slide(prs.slide_layouts[lay["index"]])
        for p in s.placeholders:
            pf = p.placeholder_format
            try:
                p.text_frame.text = f"[{pf.idx}] {pf.type}".replace("PP_PLACEHOLDER.", "")
            except Exception:
                pass
        tb = s.shapes.add_textbox(Inches(0.1), Inches(prof["slide"][1] - 0.35), Inches(6), Inches(0.3))
        tb.text_frame.text = f"layout {lay['index']}: {lay['name']}"
        r = tb.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0xD6, 0x2F, 0x2F)
        for b in prof["grid"]["forbidden"] + [d["box"] for d in lay["decor"]]:
            z = s.shapes.add_shape(1, Inches(b[0]), Inches(b[1]), Inches(max(b[2], 0.05)), Inches(max(b[3], 0.05)))
            z.fill.background(); z.line.color.rgb = RGBColor(0xD6, 0x2F, 0x2F); z.line.width = Pt(1.5)
    prs.save(out)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    src, outdir = sys.argv[1], Path(sys.argv[2])
    outdir.mkdir(parents=True, exist_ok=True)
    prof = profile(src)
    (outdir / "profile.json").write_text(json.dumps(prof, indent=1, ensure_ascii=False))
    preview(src, prof, str(outdir / "layouts_preview.pptx"))
    g = prof["grid"]
    print(f"profile: {outdir / 'profile.json'}  ({len(prof['layouts'])} layouts, slide {prof['slide']})")
    print(f"fontes: {prof['fonts']}  FOCUS={prof['palette']['FOCUS']} ACCENT={prof['palette']['ACCENT']}")
    print(f"conteúdo x {g['content_left']}–{g['content_right']}, y {g['content_top']}–{g['content_bottom']}; "
          f"colunas {g['column_x']}; linhas {g['row_y']}")
    print(f"preview: {outdir / 'layouts_preview.pptx'} → exporte para PDF e confira")
